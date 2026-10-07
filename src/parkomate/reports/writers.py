"""Render :class:`~parkomate.reports.layout.Table` objects to Excel (.xlsx) and CSV.

* Excel: one sheet per table, bold shaded header, frozen header row, column widths,
  autofilter, PASS/FAIL cells coloured (and always spelled out - never colour alone).
* CSV: one file per table, UTF-8 **with BOM** so Excel opens Marathi text correctly.

Both writers neutralise spreadsheet formula injection: text that starts with ``=``, ``+``,
``-``, ``@`` (e.g. a crafted QR code) is stored as plain text, never as a formula.
"""

from __future__ import annotations

import csv
import re
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from parkomate.reports.layout import Kind, Table, result_text

PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
PASS_FONT = Font(color="006100", bold=True)
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
FAIL_FONT = Font(color="9C0006", bold=True)
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
HEADER_FONT = Font(bold=True)

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_NUMBER_RE = re.compile(r"^-?\d+(\.\d+)?$")


def _is_injection(text: str) -> bool:
    return text.startswith(_FORMULA_START) and not _NUMBER_RE.match(text)


def cell_text(kind: Kind, value: Any, lang: str) -> str:
    """Plain-text rendering of a cell (CSV and e-mail)."""
    if kind is Kind.RESULT:
        return result_text(value, lang)
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, float):
        return f"{value:g}" if abs(value) < 1e15 else str(value)
    text = str(value)
    return f"'{text}" if _is_injection(text) else text


def safe_sheet_title(title: str, used: set[str]) -> str:
    """Excel sheet names: max 31 chars, no ``[]:*?/\\``, unique (case-insensitive)."""
    base = _BAD_SHEET_CHARS.sub("_", title).strip("'") or "Sheet"
    base = base[:31]
    candidate = base
    counter = 2
    while candidate.lower() in used:
        suffix = f" ({counter})"
        candidate = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(candidate.lower())
    return candidate


def _write_value(cell: Cell | MergedCell, kind: Kind, value: Any, lang: str) -> None:
    if isinstance(cell, MergedCell):  # pragma: no cover - writers never merge cells
        return
    if kind is Kind.RESULT:
        text = result_text(value, lang)
        cell.value = text
        if value is True:
            cell.fill = PASS_FILL
            cell.font = PASS_FONT
        elif value is False:
            cell.fill = FAIL_FILL
            cell.font = FAIL_FONT
        if text:
            cell.alignment = Alignment(horizontal="center")
        return
    if value is None:
        return
    if isinstance(value, bool):
        cell.value = str(value).lower()
        return
    if isinstance(value, int | float):
        cell.value = value
        if kind is Kind.PERCENT:
            cell.number_format = "0.0"
        return
    text = str(value)
    cell.value = text
    cell.data_type = "s"  # never a formula, whatever the text starts with


def write_table(ws: Worksheet, table: Table, lang: str, *, start_row: int = 1) -> int:
    """Write ``table`` at ``start_row``; returns the first free row after it."""
    for col, header in enumerate(table.headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for col, width in enumerate(table.widths, start=1):
        letter = get_column_letter(col)
        current = ws.column_dimensions[letter].width or 0
        ws.column_dimensions[letter].width = max(width, current)
    for offset, row in enumerate(table.rows, start=1):
        for col, (kind, value) in enumerate(zip(table.kinds, row, strict=True), start=1):
            _write_value(ws.cell(row=start_row + offset, column=col), kind, value, lang)
    return start_row + len(table.rows) + 1


def _atomic_save(workbook: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    workbook.save(tmp)
    tmp.replace(path)
    return path


def write_workbook(tables: Sequence[Table], path: Path, lang: str) -> Path:
    """One sheet per table, header row frozen."""
    workbook = Workbook()
    default = workbook.active
    if default is not None:
        workbook.remove(default)
    workbook.properties.creator = "Parkomate Station"
    used: set[str] = set()
    for table in tables:
        ws = workbook.create_sheet(safe_sheet_title(table.title, used))
        write_table(ws, table, lang)
        ws.freeze_panes = "B2" if table.freeze_first_column else "A2"
        if table.rows and table.key != "summary":
            last = get_column_letter(len(table.headers))
            ws.auto_filter.ref = f"A1:{last}{len(table.rows) + 1}"
    return _atomic_save(workbook, path)


def write_stacked_workbook(tables: Sequence[Table], path: Path, lang: str, title: str) -> Path:
    """All tables on one sheet, one under the other (individual device report)."""
    workbook = Workbook()
    ws = workbook.active
    assert ws is not None
    ws.title = safe_sheet_title(title, set())
    workbook.properties.creator = "Parkomate Station"
    row = 1
    for table in tables:
        row = write_table(ws, table, lang, start_row=row) + 1
    ws.freeze_panes = "A2"
    return _atomic_save(workbook, path)


def write_csv_tables(tables: Sequence[Table], out_dir: Path, stem: str, lang: str) -> list[Path]:
    """One ``<stem>_<table.key>.csv`` per table, UTF-8 with BOM."""
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for table in tables:
        path = out_dir / f"{stem}_{table.key}.csv"
        tmp = path.with_name(path.name + ".tmp")
        with tmp.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow([cell_text(Kind.TEXT, h, lang) for h in table.headers])
            for row in table.rows:
                writer.writerow(
                    [
                        cell_text(kind, value, lang)
                        for kind, value in zip(table.kinds, row, strict=True)
                    ]
                )
        tmp.replace(path)
        written.append(path)
    return written
