"""Reports: open the generated files and assert real cell contents."""

from __future__ import annotations

import csv
from datetime import timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from parkomate.config.settings import ReportsSettings
from parkomate.core.clock import FakeClock
from parkomate.core.enums import (
    AmbientReason,
    CheckCode,
    CounterEvent,
    IdEntryMethod,
    SessionEndReason,
    Stage,
)
from parkomate.core.errors import NotFoundError, ReportError
from parkomate.core.models import FirmwareInfo
from parkomate.data.records import ProductionRecords
from parkomate.reports.exporter import ReportExporter
from parkomate.reports.layout import SUMMARY_FIELDS, device_sheet_checks
from parkomate.reports.writers import cell_text, safe_sheet_title


@pytest.fixture
def finished_session(open_session: ProductionRecords, clock: FakeClock) -> int:
    """A session with one complete device, one rejected at B (V_C 3.21 V), one abandoned."""
    r = open_session
    r.set_session_firmware(FirmwareInfo(name="fw.bin", version="2.3.1", sha256="ab" * 32, size=9))
    r.record_ambient(27.0, AmbientReason.SESSION_START)
    good = r.start_device("24:6F:28:AA:BB:01")
    r.record_check(good, Stage.PROGRAMMING, CheckCode.A_WHITELIST, passed=True)
    r.counter_event(CounterEvent.UPLOAD_FAILURE, good)
    r.record_check(good, Stage.PROGRAMMING, CheckCode.A_UPLOAD, passed=False)
    r.counter_event(CounterEvent.UPLOAD_SUCCESS, good)
    r.record_check(good, Stage.PROGRAMMING, CheckCode.A_UPLOAD, passed=True)
    r.adjust_failure(good)
    r.record_check(good, Stage.TESTING, CheckCode.B1_COMM, passed=True)
    for i in range(1, 6):
        r.record_check(good, Stage.TESTING, CheckCode.reading(i), value=150.0 + i)
    r.record_check(
        good, Stage.TESTING, CheckCode.B3_V_C, value=3.3, limits=(3.25, 3.35), passed=True
    )
    r.set_device_id(good, "PKM-000519", IdEntryMethod.CAMERA)
    clock.advance(seconds=95)
    r.complete_device(good)

    bad = r.start_device("24:6F:28:AA:BB:02")
    r.record_check(
        bad, Stage.TESTING, CheckCode.B3_V_C, value=3.21, limits=(3.25, 3.35), passed=False
    )
    r.reject_device(
        bad,
        Stage.TESTING,
        CheckCode.B3_V_C,
        "reject.reason.out_of_range",
        {"check_key": "check.B3_V_C", "value": "3.21", "unit": "V", "low": "3.25", "high": "3.35"},
    )
    r.record_check(
        r.start_device("24:6F:28:AA:BB:03"),
        Stage.LABELING,
        CheckCode.C_QR_READ,
        value='=HYPERLINK("http://evil")',
    )
    return r.end_session(SessionEndReason.LOGOUT).id


def _exporter(r: ProductionRecords, tmp_path: Path, fmt: str = "xlsx", lang: str = "en"):
    settings = ReportsSettings(format=fmt, language=lang)  # type: ignore[arg-type]
    return ReportExporter(r.repos, tmp_path, lambda: settings, lambda: "ST-TEST")


def _rows(ws) -> list[list[object]]:  # type: ignore[no-untyped-def]
    return [list(row) for row in ws.iter_rows(values_only=True)]


def test_excel_report_sheets_and_cells(
    open_session: ProductionRecords, finished_session: int, tmp_path: Path
) -> None:
    paths = _exporter(open_session, tmp_path).export_session(finished_session)
    assert len(paths) == 1 and paths[0].suffix == ".xlsx"
    assert paths[0].name.startswith(f"session-{finished_session}_ST-TEST_")
    wb = load_workbook(paths[0])
    assert wb.sheetnames == ["Summary", "Devices", "Checks", "Rejections"]

    summary = {row[0]: row[1] for row in _rows(wb["Summary"])[1:]}
    assert summary["Operator"] == "OP1 - Omkar Operator"
    assert summary["Firmware version"] == "2.3.1"
    assert summary["Devices started"] == 3
    assert summary["Completed"] == 1
    assert summary["Rejected at B - Testing"] == 1
    assert summary["Rejected at A - Programming"] == 0
    assert summary["Abandoned"] == 1
    assert summary["Uploads failed"] == 1
    assert summary["Failures removed after a successful retry"] == 1
    assert summary["Upload failures (counted)"] == 0
    assert summary["First-pass yield (%)"] == 0.0  # the complete device had an upload failure
    assert summary["End reason"] == "Logged out"
    assert "27.0 °C (session start" in summary["Ambient readings"]
    assert wb["Summary"].freeze_panes == "B2"

    devices = wb["Devices"]
    assert devices.freeze_panes == "A2"
    header = [c.value for c in devices[1]]
    rows = _rows(devices)[1:]
    assert len(rows) == 3
    good = dict(zip(header, rows[0], strict=True))
    assert good["Device ID (QR)"] == "PKM-000519"
    assert good["Duration (s)"] == 95.0
    assert good["Upload failures"] == 1
    assert good["Point C voltage - value"] == 3.3
    assert good["Point C voltage - limits"] == "3.25 - 3.35"
    assert good["Point C voltage - result"] == "PASS"
    assert good["Sensor reading 5 - value"] == 155.0
    bad = dict(zip(header, rows[1], strict=True))
    assert bad["Status"] == "Rejected"
    assert bad["Reject stage"] == "B - Testing"
    assert bad["Failed check"] == "Point C voltage"
    assert bad["Reason"] == "Point C voltage 3.21 V - allowed 3.25-3.35 V"
    assert bad["Point C voltage - result"] == "FAIL"
    # PASS/FAIL are coloured AND spelled out.
    result_col = header.index("Point C voltage - result") + 1
    assert devices.cell(row=2, column=result_col).fill.fgColor.rgb.endswith("C6EFCE")
    assert devices.cell(row=3, column=result_col).fill.fgColor.rgb.endswith("FFC7CE")

    checks = _rows(wb["Checks"])
    assert checks[0][:6] == ["Row", "Device ID (QR)", "MAC address", "Stage", "Check code", "Check"]
    codes = [row[4] for row in checks[1:]]
    assert codes.count("A_UPLOAD") == 2  # every attempt is kept
    # Formula injection: the crafted QR text is stored as plain text.
    qr_cell = next(row for row in wb["Checks"].iter_rows(min_row=2) if row[4].value == "C_QR_READ")[
        6
    ]
    assert qr_cell.data_type == "s" and qr_cell.value.startswith("=HYPERLINK")

    rejections = _rows(wb["Rejections"])
    assert len(rejections) == 2
    rejection = dict(zip(rejections[0], rejections[1], strict=True))
    assert rejection["Box"] == "B"
    assert rejection["Value"] == 3.21
    assert rejection["Limits"] == "3.25 - 3.35"


def test_csv_report_utf8_bom_marathi(
    open_session: ProductionRecords, finished_session: int, tmp_path: Path
) -> None:
    paths = _exporter(open_session, tmp_path, fmt="csv", lang="mr").export_session(finished_session)
    assert sorted(p.name.rsplit("_", 1)[-1] for p in paths) == [
        "checks.csv",
        "devices.csv",
        "rejections.csv",
        "summary.csv",
    ]
    rejections = next(p for p in paths if p.name.endswith("_rejections.csv"))
    raw = rejections.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    with rejections.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    assert rows[0][4] == "नाकारण्याचा टप्पा"
    assert rows[1][10] == "पॉइंट C व्होल्टेज 3.21 V - परवानगी 3.25-3.35 V"
    checks = next(p for p in paths if p.name.endswith("_checks.csv"))
    text = checks.read_text(encoding="utf-8-sig")
    assert "'=HYPERLINK" in text  # neutralised in CSV too


def test_both_formats(
    open_session: ProductionRecords, finished_session: int, tmp_path: Path
) -> None:
    paths = _exporter(open_session, tmp_path, fmt="both").export_session(finished_session)
    assert len(paths) == 5


def test_device_report(
    open_session: ProductionRecords, finished_session: int, tmp_path: Path
) -> None:
    exporter = _exporter(open_session, tmp_path)
    device_row = open_session.repos.devices.list_for_session(finished_session)[0].id
    report = exporter.device_report(device_row)
    assert report["device"]["device_id"] == "PKM-000519"
    assert report["device"]["status"] == "Complete"
    assert [c["check_code"] for c in report["checks"]][:3] == [
        "A_WHITELIST",
        "A_UPLOAD",
        "A_UPLOAD",
    ]
    assert [e["event"] for e in report["counter_events"]] == [
        "upload_failure",
        "upload_success",
        "failure_adjusted",
    ]
    path = exporter.export_device(device_row)
    assert path.name == f"device_{device_row}_PKM-000519.xlsx"
    ws = load_workbook(path).active
    values = [c for row in ws.iter_rows(values_only=True) for c in row]
    assert "PKM-000519" in values and "PASS" in values


def test_range_export(
    open_session: ProductionRecords, finished_session: int, tmp_path: Path, clock: FakeClock
) -> None:
    exporter = _exporter(open_session, tmp_path)
    start = clock.now() - timedelta(days=1)
    paths = exporter.export_range(start, clock.now() + timedelta(days=1))
    wb = load_workbook(paths[0])
    assert [c.value for c in wb["Summary"][1]] == ["Field", f"Session {finished_session}"]
    with pytest.raises(ReportError):
        exporter.export_range(start - timedelta(days=10), start - timedelta(days=9))
    with pytest.raises(ReportError):
        exporter.export_sessions([])


def test_export_failures_are_typed(
    open_session: ProductionRecords, finished_session: int, tmp_path: Path
) -> None:
    exporter = _exporter(open_session, tmp_path)
    with pytest.raises(NotFoundError):
        exporter.export_session(9999)
    blocker = tmp_path / "not-a-folder"
    blocker.write_text("x", encoding="utf-8")
    with pytest.raises(ReportError):
        exporter.export_session(finished_session, out_dir=blocker)
    device_row = open_session.repos.devices.list_for_session(finished_session)[0].id
    with pytest.raises(ReportError):
        exporter.export_device(device_row, out_dir=blocker)


def test_layout_helpers() -> None:
    assert len(SUMMARY_FIELDS) > 20
    assert CheckCode.reading(3) in device_sheet_checks(3)
    assert CheckCode.reading(4) not in device_sheet_checks(3)
    used: set[str] = set()
    assert safe_sheet_title("A/B:C*?[x]", used) == "A_B_C___x_"
    assert safe_sheet_title("A/B:C*?[x]", used) == "A_B_C___x_ (2)"
    assert len(safe_sheet_title("x" * 40, set())) == 31
    from parkomate.reports.layout import Kind

    assert cell_text(Kind.RESULT, True, "en") == "PASS"
    assert cell_text(Kind.RESULT, None, "en") == ""
    assert cell_text(Kind.NUMBER, 3.30, "en") == "3.3"
    assert cell_text(Kind.TEXT, "-5", "en") == "-5"
    assert cell_text(Kind.TEXT, "+cmd", "en") == "'+cmd"
    assert cell_text(Kind.TEXT, None, "en") == ""
    assert cell_text(Kind.TEXT, True, "en") == "true"
